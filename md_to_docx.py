"""Minimal Markdown -> DOCX converter for the report files (headings, paragraphs, bold/italic/code, tables, images, lists, quotes)."""
import re, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
def runs(par, text, italic=False, size=None):
    text = text.replace("\\*", "*")
    for part in INLINE.split(text):
        if not part: continue
        if part.startswith("**"): r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"): r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9)
        elif part.startswith("*") and len(part) > 2: r = par.add_run(part[1:-1]); r.italic = True
        else: r = par.add_run(part)
        if italic: r.italic = True
        if size and not part.startswith("`"): r.font.size = Pt(size)
def convert(src, dst):
    base = Path(src).parent; d = Document(); st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    for s in d.sections: s.left_margin = s.right_margin = Inches(1); s.top_margin = s.bottom_margin = Inches(0.9)
    lines = Path(src).read_text().split("\n"); i = 0
    while i < len(lines):
        l = lines[i]
        if not l.strip() or l.strip() == "---": i += 1; continue
        m = re.match(r"(#+)\s+(.*)", l)
        if m: d.add_heading(re.sub(r"[*`]", "", m.group(2)), level=min(len(m.group(1)), 3) if len(m.group(1)) > 1 else 0); i += 1; continue
        m = re.match(r"!\[(.*)\]\((.*)\)", l)
        if m:
            p = base / m.group(2)
            if p.exists(): d.add_picture(str(p), width=Inches(5.8)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            c = d.add_paragraph(); runs(c, m.group(1), italic=True, size=9.5); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; i += 1; continue
        if l.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells): rows.append(cells)
                i += 1
            t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c, txt in enumerate(row):
                    cell = t.cell(r, c); cell.text = ""; runs(cell.paragraphs[0], txt, size=9)
                    if r == 0:
                        for rn in cell.paragraphs[0].runs: rn.bold = True
            continue
        if l.startswith(">"): p = d.add_paragraph(); runs(p, l.lstrip("> "), italic=False); p.paragraph_format.left_indent = Inches(0.3); i += 1; continue
        if re.match(r"\*Table|\*Figure", l): p = d.add_paragraph(); runs(p, l.strip("*") if l.endswith("*") and l.count("*") == 2 else l, italic=True, size=9.5); i += 1; continue
        m = re.match(r"(\d+)\.\s+(.*)", l)
        if m: p = d.add_paragraph(style="List Number"); runs(p, m.group(2)); i += 1; continue
        if re.match(r"[*-]\s+", l): p = d.add_paragraph(style="List Bullet"); runs(p, re.sub(r"^[*-]\s+", "", l)); i += 1; continue
        p = d.add_paragraph(); runs(p, l); i += 1
    d.save(dst)
if __name__ == "__main__": convert(sys.argv[1], sys.argv[2]); print("wrote", sys.argv[2])
